import datetime
import sys

import openpyxl

# Usage: python3 -I generate-openpyxl-basic.py <output.xlsx>
workbook = openpyxl.Workbook()

visible = workbook.active
visible.title = "Visible"
visible["A1"] = "Name"
visible["B1"] = "Score"
visible["C1"] = "Secret column"
visible["D1"] = "Notes"
visible["A2"] = "Alpha"
visible["B2"] = 42
visible["C2"] = "HIDDENCOLUMNTEXT"
visible["D2"] = "kept"
visible["A3"] = "HIDDENROWTEXT"
visible["B3"] = 7
visible["C3"] = "HIDDENCOLUMNTEXT"
visible["D3"] = "x"
visible["A4"] = "Total"
visible["B4"] = "=B2+B3"
visible["C4"] = "HIDDENCOLUMNTEXT"
visible["D4"] = datetime.date(2024, 1, 15)
visible["A5"] = "Flag"
visible["B5"] = True
visible["A6"] = "Multi\nline"
visible.row_dimensions[3].hidden = True
visible.column_dimensions["C"].hidden = True

secret = workbook.create_sheet("HiddenSheet")
secret["A1"] = "HIDDENSHEETTEXT"
secret.sheet_state = "hidden"

very = workbook.create_sheet("VeryHiddenSheet")
very["A1"] = "VERYHIDDENSHEETTEXT"
very.sheet_state = "veryHidden"

second = workbook.create_sheet("Second")
second["A1"] = "Beta"
second["C1"] = "gap"

visible["A1"].comment = openpyxl.comments.Comment("COMMENTTEXT", "author")

workbook.save(sys.argv[1])
