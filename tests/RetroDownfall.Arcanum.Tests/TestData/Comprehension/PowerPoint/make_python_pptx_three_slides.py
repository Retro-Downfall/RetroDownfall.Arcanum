"""Writes python-pptx-three-slides.pptx: the real-producer fixture for the PowerPoint Comprehension extractor.

Run as: python3 -I make_python_pptx_three_slides.py <output.pptx>
"""

import sys

from pptx import Presentation
from pptx.util import Inches

output = sys.argv[1]

presentation = Presentation()

# Slide 1: title and body, with speaker notes that must never be projected.
first = presentation.slides.add_slide(presentation.slide_layouts[1])
first.shapes.title.text = "Quarterly review"
first.placeholders[1].text = "Revenue grew\nCosts fell"
first.notes_slide.notes_text_frame.text = "SPEAKER NOTES ONE"

# Slide 2: hidden from the slide show (show="0"), with distinctive text.
second = presentation.slides.add_slide(presentation.slide_layouts[1])
second.shapes.title.text = "HIDDEN SLIDE TITLE"
second.placeholders[1].text = "HIDDEN SLIDE BODY"
second._element.set("show", "0")

# Slide 3: a table, a visible text box, a shape hidden through its non-visual properties, and alt text.
third = presentation.slides.add_slide(presentation.slide_layouts[5])
third.shapes.title.text = "Results"

table_shape = third.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5))
table_shape._element.xpath("./p:nvGraphicFramePr/p:cNvPr")[0].set("descr", "TABLE ALT TEXT")
table = table_shape.table
table.cell(0, 0).text = "Item"
table.cell(0, 1).text = "Count"
table.cell(1, 0).text = "Widgets"
table.cell(1, 1).text = "42"

footer = third.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1))
footer.text_frame.text = "Footer note"

hidden = third.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
hidden.text_frame.text = "HIDDEN SHAPE TEXT"
hidden._element.xpath("./p:nvSpPr/p:cNvPr")[0].set("hidden", "1")

third.notes_slide.notes_text_frame.text = "SPEAKER NOTES THREE"

presentation.save(output)
